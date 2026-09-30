from pathlib import Path

from docx import Document


def generate_docx(content: str, output_path: str) -> str:
    """
    Save document content as a DOCX file.

    Args:
        content: The document text.
        output_path: Where the DOCX file should be saved.

    Returns:
        The path to the generated file.
    """

    path = Path(output_path)

    path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    document = Document()

    for line in content.splitlines():
        line = line.strip()

        if not line:
            document.add_paragraph()
            continue

        document.add_paragraph(line)

    document.save(path)

    return str(path)